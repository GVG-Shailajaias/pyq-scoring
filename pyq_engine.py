"""
pyq_engine.py — PYQ Scoring helper script V4.0
(for PYQ_Scoring_Instructions_V4.0 and PYQ_Scoring_Rules_V4.0)

Does every MECHANICAL step of PYQ scoring, the same way every time:
loading the PYQ file, the text search, the recent window, star rating, ticker,
reference line, PYQ cells, summary table, preview page, validation and the
output xlsx.

It does NOT make judgement calls. Claude (or a person) supplies:
  - the search terms for each topic          -> Engine.search()
  - a verdict + one-line reason per candidate -> Engine.score()

State is kept in files inside a work folder, so each step can run in a
separate Python call:
  <work>/topics.json                 the day's topics (from the working sheet)
  <work>/topic_<S.No>.json           searches + candidates for one topic
  pyq_results_YYYY-MM-DD.json        the finished records AND the day's topics
                                     (single source of truth; also the
                                     "save progress" file used by resume())

Typical session:
  e = Engine(None, "2026-10-08", "work")     # finds PYQ file + rules itself
  e.load_topics_file(path)                   # once
  e.prepare("3")                             # topic text + dictionary hints
  e.search_and_show("3", ["malaria", "Plasmodium"])   # match + candidates
  e.score("3", {"2019 (Q8)": ("exact", "reason"), ...})
  e.summary_table(); e.write_preview(...)
  e.validate(); e.build_output(...); e.session_stats()

All scoring numbers are read from the rules file's ```settings block — the
script holds none of them. Change a number there, run the tests, and the
failing cases show exactly which star ratings the change moves.

Per topic, three calls: prepare() -> search_and_show() -> score().
"""

import csv
import hashlib
import json
import os
import re
import html
import shutil
from datetime import datetime

import pandas as pd

ENGINE_VERSION = "4.1"

# Name fragments used to find Project files wherever they land on disk.
# Matching ignores case, spaces, underscores, hyphens and a random upload
# prefix. The PYQ file is found whatever its version number.
FILE_PATTERNS = {
    "pyq": ("pyq", "unlockias", ".xlsx"),
    "rules": ("pyq", "scoring", "rules", ".md"),   # e.g. PYQ_Scoring_Rules.md
    "dictionary": ("synonym", "dictionary", ".csv"),
    "precedents": ("precedent", "log", ".md"),
}
SEARCH_ROOTS = ("/mnt", "/root", "/home", "/tmp", "/workspace", ".")

# ---------------------------------------------------------------- fixed format
# These belong to the output contract with the review interface, not to the
# scoring rules, so they are fixed here and not read from the rules file.
POPUP_MAX = 7             # max PYQs shown per topic
RELATED_POPUP_MAX = 5     # max related PYQs shown

SEP_TOPIC = "----------- PYQs on the Topic -----------"
SEP_RELATED = "-------- PYQs on Related Themes --------"
CELL_END = "-----------------------"
CELL_PREFIX = "UPSC CSE "
OUTPUT_HEADERS = ["S.No", "Title", "PYQ Score", "PYQ Reference", "PYQs -->"]
REQUIRED_COLS = ["ID", "Year", "Q No", "Subject", "Topic", "Difficulty",
                 "Question", "Answer Key", "Keywords"]
SEARCH_FIELDS = ["Keywords", "Question", "Solution"]
VERDICTS = ("exact", "related", "discard")

# ---------------------------------------------------------------- rules file
# Every scoring number lives in the rules file's ```settings block. The
# script has NO defaults: a missing or malformed setting stops the run.
RULE_KEYS = {
    "rules_version": str,
    "recent_window_years": int,   # "recent" = latest year back N-1 years
    "x2_recent_needed": int,      # X=2: recent exact matches needed for 3★
    "x3_recent_needed": int,      # X=3: ... for 4★
    "x4_5_recent_needed": int,    # X=4–5: ... for 4★
    "x6_plus_recent_needed": int, # X≥6: ... for 5★
    "x0_two_star_min_y": int,     # X=0: Y at or above this → 2★, else 1★
    "safety_net_max_y": int,      # safety net when X=0 and Y ≤ this
    "term_hit_limit": int,        # a term matching more questions is too broad
    "pool_limit": int,            # max candidates per topic
    "auto_word_hit_limit": int,   # automatic single-word fallbacks: max hits
    "y_warn": int,                # flag (never cut) when Y exceeds this
}


def parse_rules(text):
    """Read the ```settings block of the rules file into a dict."""
    m = re.search(r"```settings\s*\n(.*?)```", text, re.S)
    if not m:
        raise EngineError("Rules file has no ```settings block")
    raw = {}
    for line in m.group(1).splitlines():
        line = line.split("#", 1)[0].strip()
        if not line:
            continue
        if ":" not in line:
            raise EngineError(f"Bad settings line: {line!r}")
        k, v = (part.strip() for part in line.split(":", 1))
        raw[k] = v
    missing = [k for k in RULE_KEYS if k not in raw]
    unknown = [k for k in raw if k not in RULE_KEYS]
    if missing or unknown:
        raise EngineError(f"Rules settings — missing: {missing}, unknown: {unknown}")
    rules = {}
    for k, typ in RULE_KEYS.items():
        try:
            rules[k] = typ(raw[k])
        except ValueError:
            raise EngineError(f"Rules setting {k} must be {typ.__name__}: {raw[k]!r}")
    for k in RULE_KEYS:
        if RULE_KEYS[k] is int and rules[k] < 0:
            raise EngineError(f"Rules setting {k} cannot be negative")
    if rules["recent_window_years"] < 1:
        raise EngineError("recent_window_years must be at least 1")
    return rules


def load_rules(path):
    with open(path, encoding="utf-8") as f:
        return parse_rules(f.read())


# ================================================================ pure rules
# These functions take plain values and return plain values. They are what
# the tests check. The numbers come from the rules file (the `rules` dict).

def recent_start(current_year, rules):
    """First year that counts as recent (inclusive)."""
    return current_year - (rules["recent_window_years"] - 1)


def star_rating(x_years, y_years, current_year, rules):
    """Star table (Rules §3), on RAW X and Y.

    Returns (stars, rescued). stars is None when the topic is not rendered.
    """
    start = recent_start(current_year, rules)
    X, Y = len(x_years), len(y_years)
    y_recent = any(y >= start for y in y_years)

    if X == 0:
        if Y == 0:
            return None, False
        return (2 if Y >= rules["x0_two_star_min_y"] else 1), False

    if X == 1:
        if x_years[0] >= start:                # the one match is recent
            return (3 if Y >= 1 else 2), False
        return (3, True) if y_recent else (2, False)

    recent = sum(y >= start for y in x_years)
    if X == 2:
        need, base = rules["x2_recent_needed"], 3
    elif X == 3:
        need, base = rules["x3_recent_needed"], 4
    elif X in (4, 5):
        need, base = rules["x4_5_recent_needed"], 4
    else:  # X >= 6
        need, base = rules["x6_plus_recent_needed"], 5

    if recent >= need:
        return base, False
    if y_recent:
        return base, True
    return base - 1, False


def star_string(stars):
    return "★" * stars + "☆" * (5 - stars)


def sort_items(items):
    """Newest first; within a year, Q No ascending. items: dicts with year, qno."""
    return sorted(items, key=lambda it: (-it["year"], it["qno"]))


def popup_counts(X, Y):
    popup_exact = min(X, POPUP_MAX)
    remaining = POPUP_MAX - popup_exact
    popup_related = min(Y, RELATED_POPUP_MAX, remaining)
    return popup_exact, popup_related, popup_exact + popup_related


def _pyqs(n):
    return "PYQ" if n == 1 else "PYQs"


def ticker_text(X, Y):
    """Ticker (Instructions §8). Empty string when not rendered."""
    pe, pr, shown = popup_counts(X, Y)
    if X == 0 and Y == 0:
        return ""
    if X == 0:
        plus = "+" if Y > pr else ""
        return f"Related themes seen in {pr}{plus} {_pyqs(pr)}"
    if X > POPUP_MAX:
        return f"Reflected in {POPUP_MAX}+ PYQs"
    return f"Reflected in {shown} {_pyqs(shown)}"


def reference_list(shown_ids):
    """Reference list (Instructions §8): up to 3 QIDs, then '+N more'."""
    if len(shown_ids) <= 3:
        return ", ".join(shown_ids)
    return ", ".join(shown_ids[:3]) + f" +{len(shown_ids) - 3} more"


def pyq_cell(row):
    """One PYQ cell, assembled from the PYQ file row. row: dict-like."""
    ak = str(row["Answer Key"]).strip()
    ak_line = ("Answer Key: Dropped by UPSC" if ak.lower() == "dropped"
               else f"Answer Key: ({ak})")
    parts = [f"{CELL_PREFIX}{row['ID']}", row["Question"], "", ak_line]
    sol = str(row.get("Solution", "") or "").strip()
    if sol:
        parts.append(f"Solution: {sol}")
    parts.append(CELL_END)
    return "\n".join(parts)


def compose(x_items, y_items, current_year, rows_by_id, rules):
    """Everything the output needs for one topic, from its X and Y matches.

    x_items / y_items: lists of dicts {id, year, qno}. rows_by_id: QID -> row.
    """
    x_items, y_items = sort_items(x_items), sort_items(y_items)
    X, Y = len(x_items), len(y_items)
    stars, rescued = star_rating([i["year"] for i in x_items],
                                 [i["year"] for i in y_items], current_year,
                                 rules)
    if stars is None:
        return dict(stars=None, rescued=False,
                    pyq_score="Not rendered (X=0, Y=0)", pyq_reference="",
                    block=[], shown_ids=[])

    pe, pr, _ = popup_counts(X, Y)
    shown_x, shown_y = x_items[:pe], y_items[:pr]
    shown_ids = [i["id"] for i in shown_x + shown_y]

    inner = [f"X={X}", f"Y={Y}"]
    if rescued:
        inner.append("Y-rescued")
    if X:
        inner.append("X: " + ", ".join(i["id"] for i in x_items))
    if Y:
        inner.append("Y: " + ", ".join(i["id"] for i in y_items))
    pyq_score = f"{star_string(stars)} ({', '.join(inner)})"

    pyq_reference = f"{ticker_text(X, Y)} ›› {reference_list(shown_ids)}"

    block = []
    if shown_x:
        block.append(SEP_TOPIC)
        block += [pyq_cell(rows_by_id[i["id"]]) for i in shown_x]
    if shown_y:
        block.append(SEP_RELATED)
        block += [pyq_cell(rows_by_id[i["id"]]) for i in shown_y]

    return dict(stars=stars, rescued=rescued, pyq_score=pyq_score,
                pyq_reference=pyq_reference, block=block, shown_ids=shown_ids)


def safety_net_needed(X, Y, rescued, rules):
    """Rules §4 trigger: heading for the render gate, or held up by Y-rescue."""
    return (X == 0 and Y <= rules["safety_net_max_y"]) or rescued


def term_pattern(term):
    """Whole-word, case-insensitive; also matches a simple plural (s/es)."""
    words = term.strip().split()
    if not words:
        raise ValueError("empty search term")
    body = r"\s+".join(re.escape(w) for w in words)
    return re.compile(r"(?<!\w)" + body + r"(?:e?s)?(?!\w)", re.IGNORECASE)


def plain_text(s):
    """Working-sheet text with its light formatting marks removed (for reading)."""
    s = "" if s is None else str(s)
    s = s.replace("**", "").replace("__", "")
    s = re.sub(r"(?<!\w)\*(?!\s)(.+?)(?<!\s)\*(?!\w)", r"\1", s)
    return s


# ================================================================ engine
# Words never used on their own as automatic search terms.
STOPWORDS = set("""a an the and or of in on at to for from by with as is are was
were be been its it this that these those into over under about between aka also
than then which who whom whose what when where how why all any both each few more
most other some such no nor not only own same so very can will just india indian
new first second third one two three four five""".split())


def _clean_term(t):
    t = re.sub(r"\s+", " ", str(t)).strip()
    t = t.strip(" .,:;!?-–—•|/\\'\"()[]{}*_🡪→")
    return t


def baseline_terms(topic):
    """Terms taken mechanically from the topic itself: the Title (and any
    alternative name in brackets), every Tag, and every short **bold** label
    (at most 4 words) in In News / Content. Returns [(term, source)] with
    source "title", "tag" or "bold". Searched automatically on every first
    search."""
    out, seen = [], set()

    def add(t, source):
        t = _clean_term(t)
        t = re.sub(r"^(aka|a\.k\.a\.?)\s+", "", t, flags=re.I)
        t = re.sub(r"\b(19|20)\d\d\b", "", t)          # drop years like 2026
        t = _clean_term(t)
        words = t.split()
        if not words or len(t) < 3:
            return
        if all(w.lower() in STOPWORDS for w in words):
            return
        if len(words) == 1 and (len(t) < 4 and not t.isupper()):
            return
        if source == "bold" and len(words) > 4:
            return                                 # sentences, not labels
        k = t.lower()
        if k not in seen:
            seen.add(k)
            out.append((t, source))

    title = str(topic.get("Title", ""))
    for part in re.split(r"[()\[\]]", title):
        add(part, "title")
    for tag in re.split(r"[,;|/]", str(topic.get("Tags", ""))):
        add(tag, "tag")
    for field in ("In News", "Content"):
        for b in re.findall(r"\*\*(.+?)\*\*", str(topic.get(field, "")), re.S):
            add(b, "bold")
    return out


def phrase_parts(term, singles=True):
    """Shorter pieces of a multi-word term: word runs (longest first, not
    starting or ending with a stopword), then — if singles — single words of
    4+ letters."""
    words = term.split()
    if len(words) < 2:
        return []
    out = []
    for n in range(len(words) - 1, 1, -1):
        for i in range(0, len(words) - n + 1):
            run = words[i:i + n]
            if run[0].lower() in STOPWORDS or run[-1].lower() in STOPWORDS:
                continue
            out.append(_clean_term(" ".join(run)))
    for wd in (words if singles else []):
        wd = _clean_term(wd)
        if len(wd) >= 4 and wd.lower() not in STOPWORDS:
            out.append(wd)
    seen, res = set(), []
    for t in out:
        if t and t.lower() not in seen and t.lower() != term.lower():
            seen.add(t.lower())
            res.append(t)
    return res


AUTO_DISCARD_REASON = "auto-found; unrelated on preview"


class EngineError(Exception):
    pass


def _norm_name(name):
    return re.sub(r"[\s_\-]+", "", name.lower())


def _sha(path):
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def find_files(roots=None):
    """Locate the Project files by name. Returns {kind: path or None}.
    Raises if two DIFFERENT files match the same kind (identical copies are fine)."""
    roots = SEARCH_ROOTS if roots is None else roots
    found = {k: [] for k in FILE_PATTERNS}
    seen = set()
    for root in roots:
        if not os.path.isdir(root):
            continue
        for dirpath, dirnames, filenames in os.walk(root):
            dirnames[:] = [d for d in dirnames
                           if d not in ("proc", "sys", "node_modules", ".git",
                                        "site-packages", "dist-packages")]
            for fn in filenames:
                full = os.path.realpath(os.path.join(dirpath, fn))
                if full in seen:
                    continue
                n = _norm_name(fn)
                for kind, parts in FILE_PATTERNS.items():
                    if all(p.replace(" ", "") in n for p in parts[:-1]) and \
                            n.endswith(parts[-1]):
                        seen.add(full)
                        found[kind].append(full)
    out = {}
    for kind, paths in found.items():
        distinct = {}
        for p in paths:
            distinct.setdefault(_sha(p), p)
        if len(distinct) > 1:
            raise EngineError(f"More than one different '{kind}' file found: "
                              f"{sorted(distinct.values())} — keep only one in "
                              "the Project")
        out[kind] = next(iter(distinct.values()), None)
    if not out["pyq"]:
        raise EngineError("PYQ file (name containing 'PYQ' and 'UnlockIAS', "
                          ".xlsx) not found — is it uploaded to the Project?")
    return out


def load_dictionary(path):
    """Synonym dictionary CSV: Term, Also search (pipe-separated), Note."""
    entries = []
    if not path:
        return entries
    with open(path, encoding="utf-8-sig", newline="") as f:
        for r in csv.DictReader(f):
            term = (r.get("Term") or "").strip()
            if not term:
                continue
            also = [t.strip() for t in (r.get("Also search") or "").split("|")
                    if t.strip()]
            entries.append({"term": term, "also": also,
                            "note": (r.get("Note") or "").strip()})
    return entries


class Engine:
    def __init__(self, pyq_path, date, work_dir, results_path=None,
                 dictionary_path="auto", rules_path="auto"):
        """pyq_path / rules_path: a path, or None / "auto" to find it.
        date: 'YYYY-MM-DD', the date of the working sheet."""
        datetime.strptime(date, "%Y-%m-%d")
        self.date = date
        self.work = work_dir
        os.makedirs(work_dir, exist_ok=True)
        self.results_path = results_path or os.path.join(
            work_dir, f"pyq_results_{date}.json")
        files = None
        if pyq_path is None or dictionary_path == "auto" or rules_path == "auto":
            files = find_files()
        if rules_path == "auto":
            rules_path = files["rules"]
        if not rules_path:
            raise EngineError("Scoring rules file (name containing 'PYQ Scoring "
                              "Rules', .md) not found — is it in the Project?")
        self.rules_path = rules_path
        self.rules = load_rules(rules_path)
        self.pyq_path = pyq_path or files["pyq"]
        if dictionary_path == "auto":
            dictionary_path = files["dictionary"]
        self.dictionary_path = dictionary_path
        self.precedents_path = files["precedents"] if files else None
        self.dictionary = load_dictionary(dictionary_path)
        self.df = self._load_pyq(self.pyq_path)
        self.current_year = int(self.df["Year"].max())
        self.recent_start = recent_start(self.current_year, self.rules)
        self.rows = {r["ID"]: r for r in self.df.to_dict("records")}
        # if a results file already exists (same chat), restore topics from it
        if not os.path.exists(self._topics_file()) and \
                os.path.exists(self.results_path):
            self._restore_topics()

    # ---------------------------------------------------------- loading
    @staticmethod
    def _load_pyq(path):
        df = pd.read_excel(path, sheet_name="PYQ", dtype=str,
                           keep_default_na=False)
        missing = [c for c in REQUIRED_COLS if c not in df.columns]
        if missing:
            raise EngineError(f"PYQ file is missing columns: {missing}")
        if "Solution" not in df.columns:
            df["Solution"] = ""
        df["Year"] = df["Year"].astype(int)
        df["Q No"] = df["Q No"].astype(int)
        bad = df[df["ID"] != df["Year"].astype(str) + " (Q" +
                 df["Q No"].astype(str) + ")"]
        if len(bad):
            raise EngineError(f"ID does not match Year/Q No in {len(bad)} rows,"
                              f" e.g. {bad['ID'].head(3).tolist()}")
        dup = df[df["ID"].duplicated(keep=False)]
        if len(dup):
            raise EngineError(f"Duplicate IDs: {sorted(set(dup['ID']))[:5]}")
        return df

    def info(self):
        dic = (f"{len(self.dictionary)} entries" if self.dictionary_path
               else "none")
        prec = self.precedents_path or "none"
        sols = int((self.df["Solution"].str.strip() != "").sum())
        return (f"Engine v{ENGINE_VERSION}. Rules {self.rules['rules_version']}. "
                f"PYQ file: {len(self.df)} questions, "
                f"{self.df['Year'].min()}–{self.current_year}. Recent window "
                f"({self.rules['recent_window_years']} years): {self.recent_start}–"
                f"{self.current_year}. Solutions present: {sols}. Synonym "
                f"dictionary: {dic}. Precedent log: {prec}.")

    # ---------------------------------------------------------- dictionary
    def suggest_terms(self, sno):
        """Dictionary entries whose term appears in the topic's text."""
        if not self.dictionary:
            return []
        t = self.topic(sno)
        text = " ".join(plain_text(t[k]) for k in
                        ("Title", "In News", "Tags", "Content"))
        out = []
        for e in self.dictionary:
            if term_pattern(e["term"]).search(text):
                out.append(e)
        return out

    # ---------------------------------------------------------- progress
    def _restore_topics(self):
        res = self._load_results()
        if res.get("topics"):
            with open(self._topics_file(), "w", encoding="utf-8") as f:
                json.dump(res["topics"], f, ensure_ascii=False, indent=1)

    def resume(self, saved_results_path):
        """Continue from a results file saved earlier (the 'save progress' file)."""
        with open(saved_results_path, encoding="utf-8") as f:
            res = json.load(f)
        if res.get("date") != self.date:
            raise EngineError(f"Saved file is for {res.get('date')}, "
                              f"this session is {self.date}")
        if not res.get("topics"):
            raise EngineError("Saved file has no topic list — cannot resume")
        if os.path.realpath(saved_results_path) != os.path.realpath(self.results_path):
            shutil.copyfile(saved_results_path, self.results_path)
        self._restore_topics()
        done = [s for s in (t["S.No"] for t in res["topics"]) if s in res["records"]]
        todo = [s for s in (t["S.No"] for t in res["topics"]) if s not in res["records"]]
        return {"scored": done, "remaining": todo}

    def session_stats(self):
        """One line for the insights file."""
        recs = self.records()
        sn = [r for r in recs.values() if any(f.startswith("safety net:")
                                              for f in r["flags"])]
        rescued = [r for r in sn if "safety net: rescued a match" in r["flags"]]
        nr = [r for r in recs.values() if r["stars"] is None]
        return (f"Topics scored: {len(recs)} · safety net triggered {len(sn)}, "
                f"rescued a match {len(rescued)} · not rendered {len(nr)}")

    # ---------------------------------------------------------- topics
    def _topics_file(self):
        return os.path.join(self.work, "topics.json")

    def load_topics(self, rows):
        """rows: list of dicts with S.No, Title, In News, Tags, Content (Subject
        optional), in working-sheet order. S.No and Title kept exactly."""
        out, seen = [], set()
        for r in rows:
            sno = str(r["S.No"]).strip()
            if sno.endswith(".0") and sno[:-2].isdigit():
                sno = sno[:-2]
            if not sno:
                continue
            if sno in seen:
                raise EngineError(f"S.No {sno} appears twice in the working sheet")
            seen.add(sno)
            out.append({"S.No": sno, "Title": "" if r.get("Title") is None else str(r["Title"]),
                        **{k: "" if r.get(k) is None else str(r.get(k))
                           for k in ("Subject", "In News", "Tags", "Content")}})
        with open(self._topics_file(), "w", encoding="utf-8") as f:
            json.dump(out, f, ensure_ascii=False, indent=1)
        res = self._load_results()
        res["topics"] = out          # so the saved-progress file can resume
        self._save_results(res)
        return [(t["S.No"], t["Title"]) for t in out]

    def load_topics_csv(self, path):
        df = pd.read_csv(path, dtype=str, keep_default_na=False)
        return self.load_topics(df.to_dict("records"))

    def load_topics_file(self, path):
        """Working sheet saved as .xlsx (tab 'Content', else the first tab)
        or .csv. Rows without an S.No are skipped."""
        if path.lower().endswith(".csv"):
            return self.load_topics_csv(path)
        xl = pd.ExcelFile(path)
        sheet = "Content" if "Content" in xl.sheet_names else xl.sheet_names[0]
        df = pd.read_excel(path, sheet_name=sheet, dtype=str,
                           keep_default_na=False)
        df.columns = [str(c).strip() for c in df.columns]
        missing = [c for c in ("S.No", "Title") if c not in df.columns]
        if missing:
            raise EngineError(f"Working sheet tab '{sheet}' lacks columns {missing}")
        return self.load_topics(df.to_dict("records"))

    def topics(self):
        with open(self._topics_file(), encoding="utf-8") as f:
            return json.load(f)

    def topic(self, sno):
        sno = str(sno)
        for t in self.topics():
            if t["S.No"] == sno:
                return t
        raise EngineError(f"No topic with S.No {sno}")

    def topic_text(self, sno):
        t = self.topic(sno)
        return (f"S.No {t['S.No']} — {t['Title']}\nSubject: {t['Subject']}\n"
                f"In News: {plain_text(t['In News'])}\nTags: {t['Tags']}\n"
                f"Content:\n{plain_text(t['Content'])}")

    # ---------------------------------------------------------- combined calls
    def prepare(self, sno):
        """Call 1 of 3 for a topic: clean start, topic text, dictionary hints."""
        self.reset_topic(sno)
        hints = self.suggest_terms(sno)
        out = self.topic_text(sno)
        if hints:
            out += "\n\nDictionary suggestions:\n" + "\n".join(
                f"- {h['term']} -> {', '.join(h['also'])}"
                + (f"  [{h['note']}]" if h["note"] else "") for h in hints)
        return out

    def search_and_show(self, sno, terms, pass_name="stage1",
                        allow_large_reason=None):
        """Call 2 of 3: search, then (if it was accepted) the candidates."""
        res = self.search(sno, terms, pass_name, allow_large_reason)
        new = set(self._last_new_ids)
        if res["status"] == "ok" and pass_name == "stage1":
            auto = self.auto_expand(sno, terms, res["per_term"])
            new |= set(auto["new_ids"])
            res["auto_terms"] = {t: v["hits"] for t, v in auto["added"].items()}
            if auto["skipped_too_broad"]:
                res["auto_skipped_too_broad"] = auto["skipped_too_broad"]
            if auto["skipped_pool_full"]:
                res["auto_skipped_pool_full"] = auto["skipped_pool_full"]
            res["pool"] = auto["pool"]
        text = json.dumps(res, ensure_ascii=False)
        if res["status"] == "ok":
            # show only what this search added — earlier ones were already read
            text += "\n\n" + self.candidates_text(sno, only=new)
        return text

    # ---------------------------------------------------------- search
    def _work_file(self, sno):
        return os.path.join(self.work, f"topic_{sno}.json")

    def _load_work(self, sno):
        p = self._work_file(sno)
        if os.path.exists(p):
            with open(p, encoding="utf-8") as f:
                return json.load(f)
        return {"S.No": str(sno), "searches": [], "candidates": {}}

    def _save_work(self, w):
        with open(self._work_file(w["S.No"]), "w", encoding="utf-8") as f:
            json.dump(w, f, ensure_ascii=False, indent=1)

    def reset_topic(self, sno):
        """Clean candidate list for a topic (start, or a redo)."""
        self.topic(sno)
        p = self._work_file(sno)
        if os.path.exists(p):
            os.remove(p)

    def _hits(self, term):
        """{QID: set(fields)} for one term (cached)."""
        cache = self.__dict__.setdefault("_hit_cache", {})
        key = term.lower()
        if key not in cache:
            pat = term_pattern(term)
            found = {}
            for f in SEARCH_FIELDS:
                m = self.df[f].str.contains(pat)
                for qid in self.df.loc[m, "ID"]:
                    found.setdefault(qid, set()).add(f)
            cache[key] = found
        return cache[key]

    def auto_expand(self, sno, used_terms, per_term):
        """Fix 1: search the topic's own Title / Tags / bold phrases.
        Fix 2: any multi-word term with no hits is broken into shorter parts.
        A part is used only if it is narrow (multi-word: at most term_hit_limit
        hits; single word: at most auto_word_hit_limit) and the automatically
        found candidates stay within pool_limit (they are counted separately
        from Claude's own). Everything tried is recorded."""
        w = self._load_work(sno)
        R = self.rules
        used = {t.lower() for t in used_terms}
        queue = []                         # (term, origin, singles allowed)
        for t in used_terms:
            if per_term.get(t, 0) == 0:
                queue += [(p, f"part of '{t}'", True)
                          for p in phrase_parts(t, singles=len(t.split()) <= 3)]
        for t, src in baseline_terms(self.topic(sno)):
            if t.lower() not in used:
                queue.append((t, src, src in ("title", "tag")))
        added, zero, broad, over_pool, new_ids = {}, [], {}, {}, []
        tried = set(used)
        while queue:
            t, origin, singles = queue.pop(0)
            if t.lower() in tried:
                continue
            tried.add(t.lower())
            h = self._hits(t)
            n = len(h)
            limit = (R["auto_word_hit_limit"] if len(t.split()) == 1
                     else R["term_hit_limit"])
            from_topic = origin in ("title", "tag", "bold")
            if n == 0:
                zero.append(t)
                if from_topic:
                    queue += [(p, f"part of '{t}'", singles)
                              for p in phrase_parts(t, singles=singles)]
                continue
            if n > limit:
                broad[t] = n
                if from_topic and len(t.split()) > 1:
                    queue += [(p, f"part of '{t}'", singles)
                              for p in phrase_parts(t, singles=singles)]
                continue
            fresh = [q for q in h if q not in w["candidates"]]
            n_auto = sum(1 for c in w["candidates"].values() if c["found_by"] == "auto")
            if n_auto + len(fresh) > R["pool_limit"]:
                over_pool[t] = n
                continue
            for q in fresh:
                w["candidates"][q] = {"found_by": "auto", "fields": sorted(h[q]),
                                      "terms": [t]}
                new_ids.append(q)
            for q in h:
                if q not in fresh:
                    c = w["candidates"][q]
                    c["fields"] = sorted(set(c["fields"]) | h[q])
                    c["terms"] = sorted(set(c["terms"]) | {t})
            added[t] = {"hits": n, "new": len(fresh), "from": origin}
        w["searches"].append({"pass": "auto", "terms": list(added),
                              "per_term": {t: v["hits"] for t, v in added.items()},
                              "status": "ok", "override": None,
                              "skipped_broad": broad, "skipped_pool": over_pool,
                              "no_hits": zero})
        self._save_work(w)
        return {"added": added, "skipped_too_broad": broad,
                "skipped_pool_full": over_pool, "new_ids": new_ids,
                "pool": len(w["candidates"])}

    def search(self, sno, terms, pass_name="stage1", allow_large_reason=None):
        """Match terms (whole word, plural-tolerant) in Keywords, Question and
        Solution. Adds candidates to the topic only if the limits hold."""
        self.topic(sno)
        w = self._load_work(sno)
        per_term, hits = {}, {}
        for term in terms:
            th = self._hits(term)
            for qid, fields in th.items():
                hits.setdefault(qid, {"fields": set(), "terms": set()})
                hits[qid]["fields"] |= fields
                hits[qid]["terms"].add(term)
            per_term[term] = len(th)

        broad = {t: n for t, n in per_term.items() if n > self.rules["term_hit_limit"]}
        new_ids = [q for q in hits if q not in w["candidates"]]
        pool = len(w["candidates"]) + len(new_ids)
        status = "ok"
        if broad and not allow_large_reason:
            status = "narrow_terms"
        elif pool > self.rules["pool_limit"] and not allow_large_reason:
            status = "pool_too_large"

        self._last_new_ids = new_ids if status == "ok" else []
        if status == "ok":
            for qid in new_ids:
                w["candidates"][qid] = {"found_by": pass_name,
                                        "fields": sorted(hits[qid]["fields"]),
                                        "terms": sorted(hits[qid]["terms"])}
            for qid in hits:
                if qid not in new_ids:
                    c = w["candidates"][qid]
                    c["fields"] = sorted(set(c["fields"]) | hits[qid]["fields"])
                    c["terms"] = sorted(set(c["terms"]) | hits[qid]["terms"])
        w["searches"].append({"pass": pass_name, "terms": list(terms),
                              "per_term": per_term, "status": status,
                              "override": allow_large_reason})
        self._save_work(w)
        res = {"status": status, "per_term": per_term,
               "new_candidates": len(new_ids) if status == "ok" else 0,
               "pool": len(w["candidates"])}
        if broad:
            res["too_broad"] = broad
        notes = {}
        for term in terms:
            for e in self.dictionary:
                if e["note"] and (_norm_name(e["term"]) == _norm_name(term) or
                                  any(_norm_name(a) == _norm_name(term)
                                      for a in e["also"])):
                    notes[term] = e["note"]
        if notes:
            res["dictionary_notes"] = notes
        if status == "ok" and len(w["candidates"]) <= 1:
            res["hint"] = "small pool: widen once with more synonym terms"
        return res

    def add_candidate(self, sno, qid, note):
        """Add a question found by reasoning rather than by a search term."""
        if qid not in self.rows:
            raise EngineError(f"Unknown QID {qid}")
        if not note.strip():
            raise EngineError("A note saying why is required")
        w = self._load_work(sno)
        w["candidates"].setdefault(qid, {"found_by": "manual", "fields": [],
                                         "terms": [], "note": note})
        self._save_work(w)

    def _preview(self, qid, c):
        """One-line triage view of an automatically found candidate."""
        r = self.rows[qid]
        lines = [l.strip() for l in r["Question"].split("\n") if l.strip()]
        first = lines[0] if lines else ""
        hit_lines = []
        for t in c["terms"]:
            pat = term_pattern(t)
            hit_lines += [l for l in lines[1:] if pat.search(l) and l not in hit_lines]
        where = " | ".join(hit_lines[:2]) if hit_lines else (
            "matched in keywords" if "Keywords" in c["fields"] else
            "matched in solution" if "Solution" in c["fields"] else "")
        return (f"- {qid} [{r['Subject']}] ({', '.join(c['terms'])}) "
                f"{first[:160]}{' … ' + where[:160] if where else ''} "
                f"[Ans: {r['Answer Key']}]")

    def candidates_text(self, sno, only=None, full=False):
        """The candidates for reading: QID, where matched, question, answer.
        only: limit to these QIDs (e.g. the ones a widening search added)."""
        w = self._load_work(sno)
        ids = [q for q in w["candidates"] if only is None or q in only]
        out, previews = [], []
        for qid in sort_ids(ids, self.rows):
            if not full and w["candidates"][qid]["found_by"] == "auto":
                previews.append(self._preview(qid, w["candidates"][qid]))
                continue
            c = w["candidates"][qid]
            r = self.rows[qid]
            where = ", ".join(c["fields"]) or c.get("note", "manual")
            sol = r["Solution"].strip()
            out.append(f"=== {qid} [{r['Subject']} / {r['Topic']}] matched: "
                       f"{where}; terms: {', '.join(c['terms'])}\n{r['Question']}\n"
                       f"Answer Key: {r['Answer Key']}"
                       + (f"\nSolution: {sol}" if sol else ""))
        text = "\n\n".join(out)
        if previews:
            text += ("\n\n" if text else "") + (
                "Found automatically (topic terms / phrase parts) — previews; "
                "open plausible ones with e.full(sno, [QIDs]):\n" + "\n".join(previews))
        return text if text else "(no candidates)"

    def full(self, sno, qids):
        """Full text of chosen candidates (e.g. plausible auto-found ones)."""
        return self.candidates_text(sno, only=set(qids), full=True)

    # ---------------------------------------------------------- scoring
    def _load_results(self):
        if os.path.exists(self.results_path):
            with open(self.results_path, encoding="utf-8") as f:
                return json.load(f)
        return {"date": self.date, "engine_version": ENGINE_VERSION,
                "records": {}}

    def _save_results(self, res):
        with open(self.results_path, "w", encoding="utf-8") as f:
            json.dump(res, f, ensure_ascii=False, indent=1)

    def score(self, sno, verdicts, flags=(), notes="", safety_net_done=False):
        """verdicts: {QID: (verdict, one-line reason)} covering every candidate
        Claude read in full. Automatically found candidates (shown as previews)
        that are left out are recorded as discards with a standard reason.
        flags: extra flags from judgement, e.g. ['ambiguous'].
        Returns the record, or a 'needs_safety_net' notice (nothing saved)."""
        t = self.topic(sno)
        w = self._load_work(sno)
        verdicts = dict(verdicts)
        for q, c in w["candidates"].items():
            if c["found_by"] == "auto" and q not in verdicts:
                verdicts[q] = ("discard", AUTO_DISCARD_REASON)
        cands = set(w["candidates"])
        given = set(verdicts)
        if cands - given:
            raise EngineError(f"No verdict for: {sorted(cands - given)}")
        if given - cands:
            raise EngineError(f"Verdicts for non-candidates: {sorted(given - cands)}"
                              " (use add_candidate first)")
        for qid, v in verdicts.items():
            if not (isinstance(v, (list, tuple)) and len(v) == 2):
                raise EngineError(f"{qid}: give (verdict, reason)")
            if v[0] not in VERDICTS:
                raise EngineError(f"{qid}: verdict must be one of {VERDICTS}")
            if not str(v[1]).strip():
                raise EngineError(f"{qid}: a one-line reason is required")

        def items(kind):
            return [{"id": q, "year": int(self.rows[q]["Year"]),
                     "qno": int(self.rows[q]["Q No"])}
                    for q, v in verdicts.items() if v[0] == kind]

        x_items, y_items = items("exact"), items("related")
        X, Y = len(x_items), len(y_items)
        out = compose(x_items, y_items, self.current_year, self.rows, self.rules)

        if safety_net_needed(X, Y, out["rescued"], self.rules) and \
                not safety_net_done:
            w["first_pass"] = {"X": X, "Y": Y}
            self._save_work(w)
            return {"status": "needs_safety_net", "X": X, "Y": Y,
                    "rescued": out["rescued"],
                    "todo": "re-check each discard one by one, read answer keys "
                            "and solutions, search proper nouns from Content "
                            "(pass_name='safety_net'), then score(..., "
                            "safety_net_done=True)"}

        all_flags = list(flags)
        if safety_net_done:
            fp = w.get("first_pass", {"X": X, "Y": Y})
            gained = X > fp["X"] or Y > fp["Y"]
            all_flags.append("safety net: rescued a match" if gained
                             else "safety net: confirmed discard")
        if out["stars"] is None:
            all_flags.append("not rendered")
        if out["rescued"]:
            all_flags.append("Y-rescued")
        if any(str(self.rows[q]["Answer Key"]).lower() == "dropped"
               for q in out["shown_ids"]):
            all_flags.append("dropped PYQ")
        if Y > self.rules["y_warn"]:
            all_flags.append(f"Y above {self.rules['y_warn']} — check terms")
        if any(s.get("override") for s in w["searches"]):
            all_flags.append("pool limit overridden")

        record = {
            "S.No": t["S.No"], "Title": t["Title"],
            "X": X, "Y": Y,
            "x_matches": [i["id"] for i in sort_items(x_items)],
            "y_matches": [i["id"] for i in sort_items(y_items)],
            "stars": out["stars"], "rescued": out["rescued"],
            "pyq_score": out["pyq_score"], "pyq_reference": out["pyq_reference"],
            "block": out["block"], "shown_ids": out["shown_ids"],
            "flags": all_flags, "notes": notes,
            "searches": w["searches"],
            "verdicts": {q: {"verdict": v[0], "reason": v[1],
                             **w["candidates"][q]} for q, v in verdicts.items()},
            "scored_at": datetime.now().isoformat(timespec="seconds"),
            "engine_version": ENGINE_VERSION,
            "rules_version": self.rules["rules_version"],
        }
        res = self._load_results()
        before = res["records"].get(t["S.No"])
        res["records"][t["S.No"]] = record
        self._save_results(res)
        result = {"status": "saved", "stars": out["stars"], "X": X, "Y": Y,
                  "reference": out["pyq_reference"], "flags": all_flags}
        if before:
            result["change"] = diff_records(before, record)
        return result

    # ---------------------------------------------------------- review
    def records(self):
        return self._load_results()["records"]

    def summary_table(self):
        recs = self.records()
        lines = ["| S.No | Title | Stars | X | Y | PYQ Reference | Flags |",
                 "|---|---|---|---|---|---|---|"]
        for t in self.topics():
            r = recs.get(t["S.No"])
            if not r:
                lines.append(f"| {t['S.No']} | {_md(t['Title'])} | — not scored — | | | | |")
                continue
            st = star_string(r["stars"]) if r["stars"] else "Not rendered"
            lines.append(f"| {r['S.No']} | {_md(r['Title'])} | {st} | {r['X']} | "
                         f"{r['Y']} | {_md(r['pyq_reference'])} | "
                         f"{_md('; '.join(r['flags']))} |")
        return "\n".join(lines)

    def write_preview(self, path):
        recs = self.records()
        cards = []
        for t in self.topics():
            r = recs.get(t["S.No"])
            if not r:
                cards.append(f"<section class='card'><h2>{_h(t['S.No'])}. "
                             f"{_h(t['Title'])}</h2><p class='muted'>Not scored yet"
                             "</p></section>")
                continue
            if r["stars"]:
                line = (f"<p class='line'><span class='stars s{r['stars']}'>"
                        f"{star_string(r['stars'])}</span> {_h(r['pyq_reference'])}</p>")
            else:
                line = "<p class='muted'>No line shown (X=0, Y=0)</p>"
            flags = "".join(f"<span class='flag'>{_h(f)}</span>" for f in r["flags"])
            cells = "".join(
                f"<div class='sep'>{_h(c)}</div>" if c in (SEP_TOPIC, SEP_RELATED)
                else f"<pre class='q'>{_h(c)}</pre>" for c in r["block"])
            cards.append(
                f"<section class='card'><h2>{_h(r['S.No'])}. {_h(r['Title'])}</h2>"
                f"{line}<p class='meta'>X={r['X']} · Y={r['Y']} {flags}</p>"
                f"<details><summary>PYQs ({len(r['shown_ids'])})</summary>{cells}"
                f"</details></section>")
        page = PREVIEW_TEMPLATE.replace("{{DATE}}", _h(self.date)).replace(
            "{{CARDS}}", "\n".join(cards))
        with open(path, "w", encoding="utf-8") as f:
            f.write(page)
        return path

    # ---------------------------------------------------------- validation
    def validate(self):
        """Every check that must pass before the output file is built."""
        problems = []
        topics = self.topics()
        recs = self.records()
        snos = [t["S.No"] for t in topics]
        for sno in snos:
            if sno not in recs:
                problems.append(f"S.No {sno}: not scored")
        for sno in recs:
            if sno not in snos:
                problems.append(f"S.No {sno}: in results but not in the working sheet")
        titles = {t["S.No"]: t["Title"] for t in topics}
        for sno, r in recs.items():
            if sno in titles and r["Title"] != titles[sno]:
                problems.append(f"S.No {sno}: title differs from the working sheet")
            n_pyq = sum(1 for c in r["block"] if c not in (SEP_TOPIC, SEP_RELATED))
            if n_pyq > POPUP_MAX:
                problems.append(f"S.No {sno}: {n_pyq} PYQs in block (max {POPUP_MAX})")
            for q in r["x_matches"] + r["y_matches"]:
                if q not in self.rows:
                    problems.append(f"S.No {sno}: unknown QID {q}")
            if any(q not in self.rows for q in r["x_matches"] + r["y_matches"]):
                continue
            mk = lambda ids: [{"id": q, "year": int(self.rows[q]["Year"]),
                               "qno": int(self.rows[q]["Q No"])} for q in ids]
            if r.get("rules_version") != self.rules["rules_version"]:
                problems.append(f"S.No {sno}: scored under rules "
                                f"{r.get('rules_version')}, current rules are "
                                f"{self.rules['rules_version']} — redo this topic")
                continue
            again = compose(mk(r["x_matches"]), mk(r["y_matches"]),
                            self.current_year, self.rows, self.rules)
            for key in ("stars", "pyq_score", "pyq_reference", "block"):
                if again[key] != r[key]:
                    problems.append(f"S.No {sno}: {key} does not recompute "
                                    "(PYQ file or rules changed since scoring?)")
            for c in r["block"]:
                if c in (SEP_TOPIC, SEP_RELATED):
                    continue
                if not check_cell(c, self.rows):
                    problems.append(f"S.No {sno}: a PYQ cell does not match the "
                                    "PYQ file exactly")
            if r["stars"] and not r["pyq_score"].startswith(star_string(r["stars"])):
                problems.append(f"S.No {sno}: PYQ Score does not start with stars")
        return problems

    # ---------------------------------------------------------- output
    def build_output(self, out_dir):
        problems = self.validate()
        if problems:
            raise EngineError("Validation failed — no file built:\n- " +
                              "\n- ".join(problems))
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font
        recs = self.records()
        wb = Workbook()
        ws = wb.active
        ws.title = "PYQ Scores"
        ws.append(OUTPUT_HEADERS)
        for c in ws[1]:
            c.font = Font(bold=True)
        for t in self.topics():
            r = recs[t["S.No"]]
            sno = int(t["S.No"]) if t["S.No"].isdigit() else t["S.No"]
            ws.append([sno, r["Title"], r["pyq_score"], r["pyq_reference"],
                       *r["block"]])
        ws.column_dimensions["A"].width = 7
        ws.column_dimensions["B"].width = 34
        ws.column_dimensions["C"].width = 45
        ws.column_dimensions["D"].width = 45
        maxcol = max(ws.max_column, 5)
        from openpyxl.utils import get_column_letter
        for j in range(5, maxcol + 1):
            ws.column_dimensions[get_column_letter(j)].width = 60
        for row in ws.iter_rows(min_row=2):
            lines = max((str(c.value).count("\n") + 1 for c in row if c.value), default=1)
            ws.row_dimensions[row[0].row].height = min(409, max(30, 15 * lines))
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical="top")
        path = os.path.join(out_dir, f"PYQ_Scored_{self.date}.xlsx")
        wb.save(path)
        return path


# ================================================================ helpers
def sort_ids(ids, rows):
    return sorted(ids, key=lambda q: (-int(rows[q]["Year"]), int(rows[q]["Q No"])))


def check_cell(cell, rows):
    """True if a PYQ cell equals what the PYQ file says for its QID."""
    first = cell.split("\n", 1)[0]
    if not first.startswith(CELL_PREFIX):
        return False
    qid = first[len(CELL_PREFIX):]
    return qid in rows and pyq_cell(rows[qid]) == cell


def diff_records(a, b):
    """Before -> after for a redone topic."""
    def s(r):
        return star_string(r["stars"]) if r["stars"] else "Not rendered"
    return {
        "stars": f"{s(a)} -> {s(b)}",
        "X": f"{a['X']} -> {b['X']}", "Y": f"{a['Y']} -> {b['Y']}",
        "added": sorted(set(b["x_matches"] + b["y_matches"]) -
                        set(a["x_matches"] + a["y_matches"])),
        "removed": sorted(set(a["x_matches"] + a["y_matches"]) -
                          set(b["x_matches"] + b["y_matches"])),
        "moved_to_exact": sorted(set(b["x_matches"]) & set(a["y_matches"])),
        "moved_to_related": sorted(set(b["y_matches"]) & set(a["x_matches"])),
    }


def _md(s):
    return str(s).replace("|", "\\|").replace("\n", " ")


def _h(s):
    return html.escape(str(s))


PREVIEW_TEMPLATE = """<!doctype html><html lang="en"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1">
<title>PYQ Preview {{DATE}}</title><style>
:root{--bg:#f7f7f5;--card:#fff;--ink:#1d1d1b;--muted:#6b6b66;--line:#e3e2dd;
--accent:#2f5d8a;--flag:#f3ead6;--flagink:#6b4e12;--q:#fafaf8}
@media (prefers-color-scheme:dark){:root{--bg:#171716;--card:#21211f;--ink:#ecebe6;
--muted:#9c9b95;--line:#353532;--accent:#8db4dc;--flag:#3a3220;--flagink:#e6cf95;--q:#1b1b1a}}
body{margin:0;background:var(--bg);color:var(--ink);font:15px/1.5 system-ui,sans-serif}
main{max-width:860px;margin:0 auto;padding:24px 16px}
h1{font-size:20px;margin:0 0 16px}
.card{background:var(--card);border:1px solid var(--line);border-radius:10px;padding:14px 16px;margin:0 0 12px}
h2{font-size:16px;margin:0 0 6px}
.line{font-size:13px;margin:4px 0}
.stars{letter-spacing:1px}.s1,.s2{color:#9a8a5a}.s3{color:#c08a1e}.s4{color:#d06a12}.s5{color:#c0392b}
.meta{font-size:12px;color:var(--muted);margin:4px 0}
.flag{display:inline-block;background:var(--flag);color:var(--flagink);border-radius:4px;padding:0 6px;margin-left:6px}
.muted{color:var(--muted);font-size:13px}
summary{cursor:pointer;color:var(--accent);font-size:13px;margin-top:6px}
.sep{font-size:12px;color:var(--muted);margin:12px 0 4px;font-family:ui-monospace,monospace}
pre.q{white-space:pre-wrap;word-wrap:break-word;background:var(--q);border:1px solid var(--line);
border-radius:6px;padding:10px;font:13px/1.45 ui-monospace,monospace;margin:0 0 8px}
</style></head><body><main><h1>PYQ scoring preview — {{DATE}}</h1>
{{CARDS}}
</main></body></html>"""
